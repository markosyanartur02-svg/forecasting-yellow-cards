"""Fill a sheet's ACTUAL column from football-data, automatically.

    .venv/bin/python cards/fill_results.py                          # dry run
    .venv/bin/python cards/fill_results.py --write
    .venv/bin/python cards/fill_results.py --sheet "Week 3 Oct" --write

Chasing results by hand does not scale. They do not need chasing: football-data
publishes them within a few days, so the workflow is predict before kick-off,
wait, let the daily top-up fetch them, run this. The only cost is patience.

TWO RULES, BOTH OF WHICH MATTER

1. A result is written only once its card count is CONFIRMED -- at least
   `CARD_CONFIRMATION_LAG_DAYS` old. Recent rows read 0 yellows while the real
   figure is still being entered, and cards are the target, so importing one
   early corrupts the label rather than losing a feature. AFC Wimbledon v MK
   Dons really finished with 3 yellows while football-data showed 0.

2. A cell that already holds a value is NEVER overwritten. If you corrected a
   row from an official match report, that hand-entered number outranks
   anything found here.

THERE IS NO SECOND SOURCE ANY MORE

This used to accept `--source bbc`, reading results scraped from BBC Sport for
the days before football-data caught up. That scraper is retired: BBC's terms
prohibit scraping and systematic extraction, and its API host disallows all
crawlers in robots.txt -- which this project did not check until after several
hundred requests. football-data is the source the model was trained on anyway,
so the labels it gives are the right ones; BBC only ever bought speed.

Rows already filled from BBC say "actual from BBC" in Notes, and
reconcile_labels.py replaces them with football-data's figure once it is
confirmed.
"""

from __future__ import annotations

import argparse
import os
import sys

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import fixtures_week as fw
import loader

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.join(os.path.dirname(HERE), "forward_test_2026-09-18.xlsx")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--write", action="store_true",
                    help="write into the workbook; otherwise just report")
    ap.add_argument("--file", default=BOOK)
    ap.add_argument("--sheet", default="Predictions",
                    help="which sheet to fill; a weekly sheet, or Predictions")
    a = ap.parse_args()

    m = loader.flag_suspicious_cards(loader.flag_unconfirmed_cards(
        loader.load_matches(unlock_holdout=True)))
    print(f"match data runs to {m.date.max().date()}; card counts are treated "
          f"as confirmed for matches at least "
          f"{loader.CARD_CONFIRMATION_LAG_DAYS} days old\n")

    wb = load_workbook(a.file)
    ws = wb[a.sheet]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    filled = kept = pending = unpublished = review = 0
    for r in range(2, ws.max_row + 1):
        home = ws.cell(r, h["Home"]).value
        away = ws.cell(r, h["Away"]).value
        date = ws.cell(r, h["Date"]).value
        # Skip anything that is not a fixture row: the example, and the label
        # above it, which carries text in the Home column but has no date.
        if not home or not away or not date or home == "Example FC":
            continue
        date = str(date)
        existing = ws.cell(r, h["ACTUAL yellows"]).value

        hit = m[(m.HomeTeam == home) & (m.AwayTeam == away)
                & (m.date >= pd.Timestamp(date) - pd.Timedelta(days=1))
                & (m.date <= pd.Timestamp(date) + pd.Timedelta(days=2))]

        if existing is not None:
            kept += 1
            print(f"  keep    {home} v {away}: already {existing} "
                  "(hand-entered values win)")
            continue
        if len(hit) == 0:
            unpublished += 1
            continue
        row = hit.iloc[0]
        if not bool(row.cards_confirmed):
            pending += 1
            print(f"  wait    {home} v {away}: played {row.date.date()}, data "
                  f"says {int(row.yellows)} yellows but cards are not "
                  "confirmed yet")
            continue
        if bool(row.cards_suspicious):
            review += 1
            print(f"  REVIEW  {home} v {away}: data says {int(row.yellows)} "
                  f"yellows ({int(row.HY)}+{int(row.AY)}) with "
                  f"{int(row.HR)}+{int(row.AR)} red -- a team was sent off with "
                  "no yellow recorded. Fine for a straight red, impossible for "
                  "a second booking. Look it up and enter it by hand.")
            continue
        if a.write:
            ws.cell(r, h["ACTUAL yellows"], int(row.yellows))
        filled += 1
        print(f"  {'fill  ' if a.write else 'would '} {home} v {away}: "
              f"{int(row.yellows)} yellows")

    # "not in the data" is not the same as "not played": football-data runs
    # several days behind, so a fixture played on Saturday is simply absent on
    # Monday. Saying "not played yet" invites you to think the weekend never
    # happened.
    print(f"\n{filled} filled, {kept} kept as-is, {pending} awaiting "
          f"confirmation, {review} for review, {unpublished} not in the data "
          "yet")
    if a.write and filled:
        wb.save(a.file)
        print(f"saved {a.file}")
    elif not a.write:
        print("dry run -- pass --write to save")
    if pending:
        print(f"\nRe-run in a few days: football-data fills card counts in "
              f"after the result. The daily top-up will fetch them.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
