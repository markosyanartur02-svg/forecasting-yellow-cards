"""Check the unattended re-score does the right thing on the day it fires.

    .venv/bin/python cards/test_rescore.py

This code settles the forward test's labels and spends a holdout look without
anyone watching, on a date that has not arrived yet. Trusting it because it
looks right is exactly the habit this project does not have.

So it is run against copies: a copied workbook and a copied holdout log, with
the clock moved forward to the day the last labels clear the lag. Nothing here
touches the real workbook or the real log.
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import loader
import reconcile_labels as rl
import rescore_when_ready as rwr
import run_holdout
import score_forward_test as sft

PASS = FAIL = 0


def check(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name}")
    if detail:
        print(f"         {detail}")


def main() -> int:
    tmp = tempfile.mkdtemp()
    book = os.path.join(tmp, "book.xlsx")
    log = os.path.join(tmp, "holdout_log.json")
    marker = os.path.join(tmp, "rescore_done.json")
    shutil.copy(rl.BOOK, book)
    shutil.copy(run_holdout.HOLDOUT_LOG, log)
    looks_before = len(json.load(open(log)))

    real = loader.flag_unconfirmed_cards
    rl.BOOK, sft.BOOK, rwr.MARKER = book, book, marker
    run_holdout.HOLDOUT_LOG = sft.HOLDOUT_LOG = log

    print("test_declines_before_the_labels_clear")
    loader.flag_unconfirmed_cards = lambda m, **kw: real(
        m, asof=pd.Timestamp("2026-09-29"))
    rc = rwr.main()
    check("returns cleanly", rc == 0)
    check("no holdout look was spent",
          len(json.load(open(log))) == looks_before)
    check("no marker written", not os.path.exists(marker))

    print("\ntest_fires_on_the_day_they_do")
    loader.flag_unconfirmed_cards = lambda m, **kw: real(
        m, asof=pd.Timestamp("2026-09-30"))
    rc = rwr.main()
    check("returns cleanly", rc == 0)
    entries = json.load(open(log))
    check("exactly one holdout look was spent",
          len(entries) == looks_before + 1,
          f"log went from {looks_before} to {len(entries)} entries")
    if len(entries) > looks_before:
        e = entries[-1]
        check("the entry is the forward test",
              "forward" in e.get("model", "").lower(), f"{e.get('model')}")
        check("it scored all 46 fixtures", e.get("n") == 46, f"n={e.get('n')}")
    check("a marker was written", os.path.exists(marker))

    print("\ntest_never_fires_twice")
    rc = rwr.main()
    check("returns cleanly", rc == 0)
    check("still exactly one look",
          len(json.load(open(log))) == looks_before + 1,
          f"{len(json.load(open(log)))} entries")

    print("\ntest_the_hand_entered_label_survived")
    from openpyxl import load_workbook
    ws = load_workbook(book, data_only=True)["Predictions"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    wimbledon = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, h["Home"]).value == "AFC Wimbledon":
            wimbledon = ws.cell(r, h["ACTUAL yellows"]).value
    check("AFC Wimbledon is still 3, not football-data's 0",
          wimbledon == 3, f"found {wimbledon!r}")

    loader.flag_unconfirmed_cards = real
    print("\n" + "-" * 62)
    print(f"{PASS}/{PASS + FAIL} checks passed")
    print(f"(ran against copies in {tmp}; the real workbook and holdout log "
          "were not touched)")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
