"""Score the forward test once results are in. No betting prices required.

    .venv/bin/python cards/score_forward_test.py --i-am-spending-a-holdout-look

You do not need odds to find out whether the model is any GOOD. You need odds
to find out whether it is PROFITABLE. Those are different questions and the
first is answerable for free, by asking the only question that matters:

    did the model put more probability on what actually happened
    than "just use the league average" did?

That is the same metric the whole project was built on, applied to matches
that did not exist when the model was designed.

=========================================================================
PRE-REGISTERED, WRITTEN BEFORE ANY RESULT WAS SEEN (2026-09-18)
=========================================================================
Measured power on 2,036 held-out matches: at a sample of 46, a model as good
as this one beats the baseline in 75.1% of samples and reaches two sigma in
only 9.2% of them.

So this test is badly underpowered, and that is not a flaw to be apologised
for afterwards -- it is the reason the following is fixed in advance:

  model gain > 0, 2 sigma or better  -> STRONG (expected only 9% of the time
                                        even if the model is good; treat a
                                        result this clean with mild suspicion)
  model gain > 0, below 2 sigma      -> CONSISTENT. The expected outcome.
                                        Not proof. Says "nothing is broken".
  model gain <= 0, above -2 sigma    -> INCONCLUSIVE, leaning bad. Happens
                                        ~25% of the time to a good model.
  model gain <= 0, beyond -2 sigma   -> WARNING. Rare by luck; suspect a real
                                        problem, most likely in the features
                                        for a season the model has barely seen.

Nothing here is a verdict on the model. One weekend is a smoke test. The
verdict needs roughly 1,000 matches, which is half a season across the four
divisions.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime, timezone

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import features
import loader
import models
import freeze as freeze_mod
from run_holdout import HOLDOUT_LOG, log_look, looks_taken

BOOK = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "forward_test_2026-09-18.xlsx")


def read_results(path: str) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True)
    ws = wb["Predictions"]
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    rows = []
    for r in range(2, ws.max_row + 1):
        home = ws.cell(r, head["Home"]).value
        if not home or home == "Example FC":
            continue
        actual = ws.cell(r, head["ACTUAL yellows"]).value
        rows.append({
            "date": ws.cell(r, head["Date"]).value,
            "division": ws.cell(r, head["Div"]).value,
            "home": home, "away": ws.cell(r, head["Away"]).value,
            "referee": ws.cell(r, head["Referee used"]).value,
            "baseline": ws.cell(r, head["League avg"]).value,
            "mu": ws.cell(r, head["Model expects"]).value,
            "pmf": [ws.cell(r, head[f"P({k})"]).value for k in range(9)],
            "actual": actual,
        })
    return pd.DataFrame(rows)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--i-am-spending-a-holdout-look", action="store_true",
                    dest="spend")
    ap.add_argument("--file", default=BOOK)
    ap.add_argument("--dry-run", action="store_true",
                    help="score without recording a look; for testing the "
                         "scorer itself against simulated results")
    a = ap.parse_args()

    df = read_results(a.file)
    done = df[df.actual.notna()].copy()
    print(f"{len(df)} fixtures in the workbook, {len(done)} with a result\n")
    if len(done) == 0:
        print("No results entered yet. Fill in the ACTUAL yellows column.")
        return 1
    if not a.spend:
        print(f"This scores holdout data ({looks_taken()} looks spent on the "
              "2026/27 holdout so far).\n"
              "Pass --i-am-spending-a-holdout-look to proceed.")
        return 1

    # The accumulating forward test is one experiment only while the model
    # specification holds still. Checked rather than remembered.
    print("checking the model is still frozen ...")
    if not freeze_mod.do_verify():
        print("\nREFUSING TO SCORE.\n"
              "The specification has changed since the freeze. Weeks scored\n"
              "before the change measured a different model, so adding this\n"
              "week to them measures nothing.\n\n"
              "If the change was cosmetic -- a docstring, a comment -- re-freeze\n"
              "with cards/freeze.py --freeze \"reason\" and run again.\n"
              "If it was real, the honest move is to start a fresh forward test\n"
              "for the new model and keep the two results apart.")
        return 1
    print()

    # ---- the model's dispersion, and a baseline fitted WITHOUT these matches
    m = loader.load_matches()                     # development only
    feat = features.build(m)
    cols = features.feature_columns(feat)
    glm = models.NBGLM(cols).fit(feat)
    base = models.LeagueSeasonNB().fit(feat)
    alpha_g, alpha_b = glm.alpha_, base.alpha_

    y = done.actual.to_numpy(float)
    # model: read the pre-registered distribution straight from the sheet, so
    # the score uses exactly what was committed to before kick-off
    pm = np.array([r[int(min(v, 8))] for r, v in zip(done.pmf, y)], dtype=float)
    ll_model = np.log(np.clip(pm, 1e-12, None))
    ll_base = models.nb_logpmf(y, done.baseline.to_numpy(float), alpha_b)

    d = ll_model - ll_base
    ok = np.isfinite(d)
    d, y = d[ok], y[ok]
    n = len(d)
    mean, sd = float(d.mean()), float(d.std(ddof=1))
    se = sd / np.sqrt(n) if n > 1 else np.nan
    sigma = mean / se if se and se > 0 else np.nan

    print("=" * 66)
    print(f"FORWARD TEST RESULT   ({n} matches)")
    print("=" * 66)
    print(f"  mean yellows predicted : {done[ok].mu.mean():.2f}")
    print(f"  mean yellows actual    : {y.mean():.2f}")
    print(f"  bias                   : {y.mean() - done[ok].mu.mean():+.2f}")
    print(f"  mean absolute error    : {np.abs(y - done[ok].mu.to_numpy()).mean():.2f}")
    print(f"\n  model log-likelihood gain over the league average:")
    print(f"     {mean:+.5f} per match   sd {sd:.4f}   -> {sigma:+.1f} sigma")
    print(f"     model better in {(d > 0).sum()}/{n} matches")

    if mean > 0 and sigma > 2:
        verdict, note = "STRONG", ("expected in only ~9% of samples this size "
                                   "even for a good model -- pleasing, but "
                                   "treat it with mild suspicion")
    elif mean > 0:
        verdict, note = "CONSISTENT", ("the expected outcome. Not proof. It "
                                       "says nothing is broken")
    elif sigma > -2:
        verdict, note = "INCONCLUSIVE", ("leaning bad, but this happens to a "
                                         "good model ~25% of the time at n=46")
    else:
        verdict, note = "WARNING", ("rare by luck -- suspect a real problem")
    print(f"\n  VERDICT: {verdict} -- {note}")

    # ---- calibration, the other price-free check
    print("\n  CALIBRATION")
    for line in (2.5, 3.5, 4.5):
        k = int(np.floor(line))
        pred = np.array([sum(r[k + 1:]) for r in done[ok].pmf])
        print(f"     over {line}: predicted {pred.mean():.1%}   "
              f"actual {(y > line).mean():.1%}   "
              f"gap {pred.mean() - (y > line).mean():+.1%}")

    # ---- where the gain came from
    dd = done[ok].copy()
    dd["gain"] = d
    no_ref = dd.referee.isna() | (dd.referee == "not known")
    print(f"\n  BY DIVISION   ({int(no_ref.sum())} of {len(dd)} fixtures had no "
          "referee appointed)")
    for div in ("E0", "E1", "E2", "E3"):
        g = dd[dd.division == div]
        if len(g) == 0:
            continue
        print(f"     {loader.DIVISIONS[div]:<16} n={len(g):>3}  "
              f"gain {g.gain.mean():+.5f}  better in {(g.gain > 0).sum()}/{len(g)}"
              f"   predicted {g.mu.mean():.2f} actual {g.actual.mean():.2f}")

    if a.dry_run:
        print(f"\n{'=' * 66}")
        print("DRY RUN -- no holdout look recorded.")
        return 0

    log_look({
        "model": "NBGLM forward test", "n": int(n),
        "holdout_definition": list(loader.HOLDOUT_SEASONS),
        "mean_gain_vs_baseline": mean, "sigma": float(sigma),
        "verdict": verdict,
        "note": "pre-registered forward test on 46 fixtures, 18-20 Sep 2026; "
                "predictions fixed before kick-off",
    })
    print(f"\n{'=' * 66}")
    print(f"holdout looks spent on 2026/27: {looks_taken()}")
    print(f"  (and {looks_taken(current_only=False)} in the project's history, "
          "counting the retired 2024/25+2025/26 holdout)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
