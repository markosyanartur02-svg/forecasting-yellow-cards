"""Watch the EFL for referee appointments, and say when they land.

    .venv/bin/python cards/check_referees.py            # check, quietly
    .venv/bin/python cards/check_referees.py --verbose  # print what was parsed

Run from the daily launchd job. Most days it prints one line and exits. On the
day the appointments appear it says so, saves them, and names the fixtures it
can fill in.

The referee is 48% of the model's value, so a week's predictions made before
the appointments are the weaker version of the model. Knowing the hour they
land is worth a daily HTTP request.

THE MARKUP IS NOT WORTH PARSING

Four layouts in four weeks:

    <h2>Friday, 18 September</h2> <h3>Sky Bet Championship</h3>
    <h4>Bristol City v Watford (20:00)</h4> <p>Referee: Oliver Langford ...

    <h3>Saturday, 26 September</h3> <p><strong>Sky Bet League One</strong></p>
    <p>Cambridge United v AFC Wimbledon (15:00)<br />BEN TONER<br />...

    <p>Accrington Stanley v Sunderland U21 (19:00)<br />STEVEN COPELAND<br />

    <h3>Friday, 9 October</h3> <h4>EFL Championship</h4>
    <p><strong>West Ham United v QPR (20:00)</strong><br />Referee: Ruebyn ...

Chasing each one cost a week of referees twice: the fixtures parsed, or did
not, and the job reported "0 fixtures, still waiting" on days when the
appointments were published and sitting there.

So this does not parse markup. It flattens the article to lines -- <br>, </p>
and every closing heading become line breaks -- and then reads the lines the
way a person does, in order:

    a line naming a competition  -> that is the division from here on
    a line naming a month        -> that is the date from here on
    "Home v Away (HH:MM)"        -> a fixture
    the next "Referee: X", or
      the next ALL-CAPS NAME     -> that fixture's referee

Nothing about which tag holds what. A fifth layout that keeps the same reading
order needs no change here.

The content is inside the page's Nuxt payload rather than its rendered body,
so this reads the raw HTML rather than stripped text.

WHAT COUNTS AS PENDING

The fixtures still waiting for a referee are read from the workbook itself --
every "Week ..." sheet, every row dated today or later whose Referee column
says "not known". Not from another module's constant.

That coupling has now failed twice. First the parser dropped a layout and
reported zero fixtures; then run_week_sheet stopped having a FIXTURES list at
all when it moved to date ranges, and this asked an empty list whether it
contained the 31 appointments that had just been published. Both times the
answer was "still waiting" on a day when the referees were sitting there.

The workbook is the thing that actually knows what is pending, so ask it.

VALIDATED AGAINST KNOWN GROUND TRUTH

The 18-20 September article is parsed and checked against the 45 referees that
were entered by hand into fixtures_week.py last week. All 45 agree, which is
the only reason this is trusted to fill them in automatically.
"""

from __future__ import annotations

import argparse
import html as htmlmod
import json
import os
import re
import sys
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import fetch_results_bbc as bbc
import net

BOOK = os.path.join(os.path.dirname(HERE), "forward_test_2026-09-18.xlsx")

CATEGORY = "https://www.efl.com/category/match-officials/"
STATE = os.path.join(HERE, "referee_appointments.json")

COMPETITIONS = {
    "championship": "E1",
    "league one": "E2",
    "league two": "E3",
}
# Articles that are not a league round. Kept separate rather than ignored, so
# "appointments are out" never fires on a cup article.
CUPS = ("trophy", "carabao", "play-off", "playoff", "laws-of-the-game", "laws")

LINK = re.compile(r'"(/news/(\d{4})/([a-z]+)/(\d{2})/([a-z0-9\-]*referee[a-z0-9\-]*)/)"')

# A fixture line: "Home v Away (19:45)". Team names hold no digits-colon
# pair and no "v" of their own, so this is unambiguous once the markup is gone.
FIXTURE_LINE = re.compile(
    r"^(.{2,40}?)\s+v\s+(.{2,40}?)\s*\((\d{1,2}[:.]\d{2})\)\s*$")
# "Referee: Oliver Langford Assistants: George Byrne ... Fourth Official:
# Simon Hooper" arrives as ONE line in some layouts, so the name has to stop
# at the next role. Without that, the surname taken was the fourth official's
# and every referee in the 18-20 September article came out wrong while
# looking perfectly plausible.
REFEREE_LINE = re.compile(
    r"^Referee\s*:\s*(.+?)\s*(?:\b(?:Assistants?|Fourth Official|VAR|AVAR|"
    r"Observer)\b.*)?$", re.I)
# A bare personal name on its own line, which is how the referee appears when
# the article does not label it: "BEN TONER" in the cup articles, "Dean
# Whitestone" in the League One and Two sections. Two to four capitalised
# words and nothing else.
#
# The anchors do the important work. The assistants sit on the very next line,
# unlabelled too -- "Bradley Hall and Matt Archibald" -- and are rejected
# because "and" is lower case, so the match cannot reach the end of the line.
# The referee is simply the first such line after the fixture.
NAME_LINE = re.compile(r"^([A-Z][A-Za-z'\-]+(?:\s+[A-Z][A-Za-z'\-]+){1,3})$")
# "Assistants:", "Fourth Official:" and the like are not referees.
NOT_REFEREE = re.compile(r"^(assistants?|fourth official|var|avar|observer)\b",
                         re.I)

MONTHS = ("january", "february", "march", "april", "may", "june", "july",
          "august", "september", "october", "november", "december")


def clean(s: str) -> str:
    return " ".join(htmlmod.unescape(re.sub(r"<[^>]+>", "", s)).split())


def model_name(full: str) -> str:
    """'Oliver Langford' -> 'o langford', the form the model's data uses."""
    parts = clean(full).split()
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0].lower()
    return f"{parts[0][0].lower()} {parts[-1].lower()}"


def list_articles(raw: str) -> list[dict]:
    out, seen = [], set()
    for path, year, month, day, slug in LINK.findall(raw):
        if path in seen:
            continue
        seen.add(path)
        out.append({
            "url": "https://www.efl.com" + path,
            "slug": slug,
            "is_cup": any(c in slug for c in CUPS),
            "published": f"{year}-{month}-{day}",
        })
    return out


# The article body inside the Nuxt payload. Reading the whole page instead
# picks up the site navigation, whose menu says "Championship", "League One",
# "League Two" -- which made the EFL Trophy article report 22 Championship
# fixtures. The body has to be isolated before anything else is read.
OPENER = re.compile(r"content:\s*(['\"])")


def article_body(raw: str) -> str:
    """Every `content:` string in the payload, scanned not regexed.

    The payload uses single quotes in some articles and double in others, and
    the strings carry escaped quotes and backslash-n. A regex for one style
    silently returned nothing for the other -- which is how 18-20 September
    went from parsing correctly to parsing not at all. Scanning handles both
    without caring.
    """
    chunks = []
    for m in OPENER.finditer(raw):
        q = m.group(1)
        i = m.end()
        buf = []
        while i < len(raw):
            c = raw[i]
            if c == "\\" and i + 1 < len(raw):
                buf.append(raw[i:i + 2])
                i += 2
                continue
            if c == q:
                break
            buf.append(c)
            i += 1
        chunks.append("".join(buf))
    keep = [c for c in chunks if " v " in c and "(" in c]
    return "\n".join(keep or chunks)


def to_lines(raw: str) -> list[str]:
    """The article as the lines a reader would see, markup removed.

    Not stripping <script> first, which is the obvious move and deletes
    everything: the content lives inside the page's Nuxt payload, so the
    script block IS the text.
    """
    t = article_body(raw)
    # Block boundaries become line breaks. \\n is how the payload spells a
    # newline inside its JS string, so it counts as one too.
    t = re.sub(r"(?i)<br\s*/?>", "\n", t)
    t = re.sub(r"(?i)</(p|h[1-6]|li|div|tr|td)\s*>", "\n", t)
    t = t.replace("\\n", "\n")
    t = re.sub(r"<[^>]+>", "", t)
    out = []
    for line in t.split("\n"):
        line = clean(line)
        if line:
            out.append(line)
    return out


def classify(text: str) -> tuple[str, str] | None:
    """A marker's meaning comes from its words, not from its tag."""
    low = text.lower()
    for needle in COMPETITIONS:
        if needle in low:
            return ("comp", text)
    if any(mon in low for mon in MONTHS):
        return ("date", text)
    return None


def parse_article(raw: str) -> list[dict]:
    """Fixtures with date, competition and referee, read in reading order."""
    lines = to_lines(raw)
    date = comp = None
    out = []
    i = 0
    while i < len(lines):
        line = lines[i]
        fx = FIXTURE_LINE.match(line)
        if not fx:
            got = classify(line)
            if got:
                if got[0] == "comp":
                    comp = got[1]
                else:
                    date = got[1]
            i += 1
            continue

        # The referee is the next line that says so, or the next bare name in
        # capitals -- whichever comes first, within a few lines. Looking
        # further would start picking up the following fixture's official.
        ref = None
        for j in range(i + 1, min(i + 5, len(lines))):
            nxt = lines[j]
            if FIXTURE_LINE.match(nxt) or classify(nxt):
                break
            m = REFEREE_LINE.match(nxt)
            if m:
                ref = m.group(1)
                break
            if NOT_REFEREE.match(nxt):
                continue
            m = NAME_LINE.match(nxt)
            if m:
                ref = m.group(1)
                break
        if ref:
            div = None
            if comp:
                low = comp.lower()
                for needle, code in COMPETITIONS.items():
                    if needle in low:
                        div = code
                        break
            out.append({"date_text": date, "competition": comp,
                        "division": div, "home": fx.group(1),
                        "away": fx.group(2), "time": fx.group(3),
                        "referee_full": ref, "referee": model_name(ref)})
        i += 1
    return out


def pending_fixtures(path: str = BOOK) -> set:
    """(home, away) in football-data spelling, not yet played, no referee."""
    from datetime import date
    from openpyxl import load_workbook
    if not os.path.exists(path):
        return set()
    wb = load_workbook(path, data_only=True)
    today = date.today().isoformat()
    out = set()
    for name in wb.sheetnames:
        if not name.startswith("Week"):
            continue
        ws = wb[name]
        h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        if not all(k in h for k in ("Date", "Home", "Away", "Referee")):
            continue
        for r in range(2, ws.max_row + 1):
            d = ws.cell(r, h["Date"]).value
            home = ws.cell(r, h["Home"]).value
            away = ws.cell(r, h["Away"]).value
            ref = ws.cell(r, h["Referee"]).value
            if not d or not home or not away:
                continue
            if str(d) < today:
                continue
            if ref and str(ref).strip().lower() not in ("", "not known"):
                continue
            out.add((str(home), str(away)))
    return out


def load_state() -> dict:
    if os.path.exists(STATE):
        with open(STATE) as f:
            return json.load(f)
    return {}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--verbose", action="store_true")
    ap.add_argument("--refresh", action="store_true",
                    help="re-parse the newest article even if it was already "
                         "seen; use after fixing the parser")
    ap.add_argument("--url", help="parse this article instead of looking one up")
    a = ap.parse_args()

    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    if a.url:
        fx = parse_article(bbc.get(a.url))
        print(f"{stamp} parsed {len(fx)} fixtures from {a.url}")
        for f in fx:
            print(f"   {f['division'] or '--'} {f['date_text']}  "
                  f"{f['home']} v {f['away']} ({f['time']})  "
                  f"{f['referee_full']}  ->  {f['referee']}")
        return 0

    if not net.wait_for_network(
            log=lambda m: print(f"{stamp} referees: {m}")):
        print(f"{stamp} referees: skipped, no network")
        return 0

    try:
        arts = list_articles(bbc.get(CATEGORY))
    except Exception as e:
        # A scheduled job that dies with a traceback looks identical in the
        # log to one that never ran. Say what happened and exit cleanly.
        print(f"{stamp} referees: could not reach the EFL ({type(e).__name__}) "
              "-- will try again tomorrow")
        return 0
    league = [x for x in arts if not x["is_cup"]]
    if not league:
        print(f"{stamp} referees: no league appointment article found")
        return 0
    newest = league[0]

    state = load_state()
    if state.get("url") == newest["url"] and not a.refresh:
        print(f"{stamp} referees: nothing new "
              f"(latest is still '{newest['slug']}')")
        return 0

    fx = parse_article(bbc.get(newest["url"]))
    covered = [f for f in fx if f["division"]]

    # "New article" is not the same as "this week's appointments". On a first
    # run the newest article is often last week's, and announcing that would
    # be a false alarm every time the state file is missing. What matters is
    # whether it covers the round we are waiting on.
    pending, hit = set(), set()
    try:
        pending = pending_fixtures()
        # The article writes full club names, the workbook uses
        # football-data's short ones, so resolve before comparing.
        names = sorted({n for fx in pending for n in fx})
        got = set()
        for f in covered:
            h2 = bbc.resolve(f["home"], names)
            a2 = bbc.resolve(f["away"], names)
            if h2 and a2:
                got.add((h2, a2))
        hit = pending & got
    except Exception as e:
        print(f"{stamp} referees: could not read the workbook "
              f"({type(e).__name__}: {e})")

    record = {
        "url": newest["url"], "slug": newest["slug"],
        "published": newest["published"],
        "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "n_fixtures": len(covered),
        "fixtures": covered,
    }
    with open(STATE, "w") as f:
        json.dump(record, f, indent=2)

    if hit:
        print(f"{stamp} REFEREE APPOINTMENTS ARE OUT: {newest['slug']}")
        print(f"{stamp}   {newest['url']}")
        print(f"{stamp}   {len(covered)} league fixtures parsed, covering "
              f"{len(hit)} of the {len(pending)} in the pending sheet")
        for h, aw in sorted(pending - hit):
            print(f"{stamp}      still without one: {h} v {aw}")
        print(f"{stamp}   re-run: .venv/bin/python cards/run_week_sheet.py")
    else:
        print(f"{stamp} referees: new article '{newest['slug']}' "
              f"({len(covered)} fixtures), but none of them are the "
              f"{len(pending)} in the pending sheet -- still waiting")

    if a.verbose:
        for f in covered:
            print(f"   {f['division']} {f['home']} v {f['away']}: "
                  f"{f['referee_full']} -> {f['referee']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
