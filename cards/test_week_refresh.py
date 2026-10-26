"""Refreshing a week's sheet must never lose or rewrite a played row.

    .venv/bin/python cards/test_week_refresh.py

auto_week_sheet.py re-runs run_week_sheet.py twice a day on the same sheet.
Three things have to hold for that to be safe:

1. A played match keeps its row even after fixtures.csv stops listing it,
   which happens as soon as football-data posts the result. Before this was
   handled, a refresh after an update deleted the weekend's rows, prices and
   all.
2. A kept row that ends up on a different row keeps formulas that point at
   itself. Copying "=...$AV2..." verbatim onto row 3 reads the wrong match.
3. A match that kicked off today is not re-predicted, even though its date is
   not yet in the past.
"""

from __future__ import annotations

import datetime as dt
import os
import sys
import tempfile

import numpy as np
from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import run_week_sheet as rws

PASS = FAIL = 0
SHEET = "Week refresh"


def check(name: str, got, want, tol: float | None = None) -> None:
    global PASS, FAIL
    if tol is not None and isinstance(got, (int, float)) \
            and isinstance(want, (int, float)):
        ok = abs(float(got) - float(want)) <= tol
    else:
        ok = got == want
    if ok:
        PASS += 1
        print(f"  [PASS] {name}: {got}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name}: got {got!r}, expected {want!r}")


def pred(date, div, home, away, mu, over35=0.4151):
    pmf = np.zeros(9)
    pmf[3] = 1.0
    return {"division": div, "date": date, "home": home, "away": away,
            "referee": None, "mu": mu, "baseline": 3.9, "pmf": pmf,
            "over": [0.9, 0.8, over35, 0.2, 0.1, 0.05], "ref_n": 0,
            "ref_rel": None, "mode": 3, "iv": "2-5"}


def rows_by_match(path):
    ws = load_workbook(path)[SHEET]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    out = {}
    for r in range(2, ws.max_row + 1):
        home, away = ws.cell(r, h["Home"]).value, ws.cell(r, h["Away"]).value
        if home and away:
            out[(home, away)] = (r, {name: ws.cell(r, c).value
                                     for name, c in h.items()})
    return out, h


def test_kicked_off():
    print("\nkick-off is judged by date and UK time")
    now = dt.datetime(2026, 10, 24, 16, 0, tzinfo=rws.UK)
    check("yesterday has started", rws.kicked_off("2026-10-23", "20:00", now), True)
    check("tomorrow has not", rws.kicked_off("2026-10-25", "12:00", now), False)
    check("today 15:00 has, at 16:00", rws.kicked_off("2026-10-24", "15:00", now), True)
    check("today 17:30 has not, at 16:00", rws.kicked_off("2026-10-24", "17:30", now), False)
    check("today with no time counts as started",
          rws.kicked_off("2026-10-24", "", now), True)


def test_refresh_keeps_played_rows():
    print("\na refresh keeps a played row that fixtures.csv no longer lists")
    today = rws.uk_now().date()
    yesterday = (today - dt.timedelta(days=1)).isoformat()
    tomorrow = (today + dt.timedelta(days=1)).isoformat()
    tmp = os.path.join(tempfile.mkdtemp(), "refresh.xlsx")
    Workbook().save(tmp)

    rws.write(tmp, SHEET, [pred(yesterday, "E3", "A", "B", 3.3),
                           pred(tomorrow, "E3", "C", "D", 3.0)], "fp")
    wb = load_workbook(tmp)
    ws = wb[SHEET]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    for name, v in (("Settles on", "yellow cards"), ("Card line", 3.5),
                    ("Over price", 1.80), ("Under price", 2.05),
                    ("ACTUAL yellows", 2), ("Notes", "my price")):
        ws.cell(2, h[name], v)
    wb.save(tmp)

    # Second run: A v B has gone from fixtures.csv, C v D has a new forecast,
    # and Q v R sorts in front of A v B, pushing it down a row. (main() would
    # never create Q v R after kick-off; write() is tested on its own here.)
    rws.write(tmp, SHEET, [pred(tomorrow, "E3", "C", "D", 3.7),
                           pred(yesterday, "E2", "Q", "R", 4.0)], "fp")
    rows, h = rows_by_match(tmp)
    check("played row still there", ("A", "B") in rows, True)
    r, ab = rows[("A", "B")]
    check("it moved down a row", r, 3)
    check("its price survived", ab["Over price"], 1.80)
    check("its note survived", ab["Notes"], "my price")
    check("its forecast was not rewritten", ab["Model expects"], 3.3)
    check("the unplayed row was refreshed", rows[("C", "D")][1]["Model expects"], 3.7)

    import formulas
    sol = formulas.ExcelModel().loads(tmp).finish().calculate()
    base = f"'[{os.path.basename(tmp)}]{SHEET.upper()}'!"
    letter = lambda name: rws.get_column_letter(h[name])

    def g(name, row):
        v = sol[f"{base}{letter(name)}{row}"].value[0, 0]
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v

    # the same numbers test_week_sheet.py checks for this price, on row 3 now
    check("moved row's edge reads its own prices", g("Edge", 3), 0.4151 - 0.53253, 1e-4)
    check("moved row's verdict", g("BET?", 3), "UNDER")
    check("moved row's result", g("Won?", 3), "WON")


def test_started_today_not_rewritten():
    print("\na match that kicked off today keeps its morning forecast")
    today = rws.uk_now().date().isoformat()
    tmp = os.path.join(tempfile.mkdtemp(), "today.xlsx")
    Workbook().save(tmp)
    rws.write(tmp, SHEET, [pred(today, "E1", "C", "D", 3.0)], "fp")
    rws.write(tmp, SHEET, [pred(today, "E1", "C", "D", 9.9)], "fp",
              frozenset({(today, "C", "D")}))
    rows, _ = rows_by_match(tmp)
    check("forecast kept from before kick-off", rows[("C", "D")][1]["Model expects"], 3.0)


def test_auto_week():
    print("\nthe unattended run: one sheet per Thursday-to-Wednesday week")
    import auto_week_sheet as auto
    for day, want in (("2026-10-26", ("2026-10-22", "2026-10-28")),    # Monday
                      ("2026-10-28", ("2026-10-22", "2026-10-28")),    # Wednesday
                      ("2026-10-29", ("2026-10-29", "2026-11-04")),    # Thursday
                      ("2026-11-01", ("2026-10-29", "2026-11-04"))):   # Sunday
        s, e = auto.week_window(dt.date.fromisoformat(day))
        check(f"week containing {day}", (s.isoformat(), e.isoformat()), want)

    folder = tempfile.mkdtemp()
    book = os.path.join(folder, "book.xlsx")
    Workbook().save(book)
    check("closed workbook is not 'open'", auto.open_in_excel(book), False)
    open(os.path.join(folder, "~$book.xlsx"), "w").close()
    check("Excel's lock file means open", auto.open_in_excel(book), True)


def main() -> int:
    test_kicked_off()
    test_refresh_keeps_played_rows()
    test_started_today_not_rewritten()
    test_auto_week()
    print("\n" + "-" * 62)
    print(f"{PASS}/{PASS + FAIL} checks passed")
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
