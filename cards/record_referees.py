"""Record who actually refereed a fixture, without rewriting the prediction.

    .venv/bin/python cards/record_referees.py --sheet "Week 9-12 Oct" \
        --url https://www.efl.com/news/2026/october/07/referee-appointments--9---15-october/

Writes a separate "Referee (actual)" column. The "Referee" column is left
exactly as it was.

WHY TWO COLUMNS

"Referee" records what the MODEL KNEW when it made the prediction. For the
26 September and 9-12 October rounds that is "not known", because the watcher
failed to notice the appointments until after kick-off. Writing the real name
into that column would claim a prediction used information it never had, which
is the same hindsight this project spends most of its guards preventing.

So the truth goes in its own column, where it is useful without being a lie:
it lets next season's model be checked against the referees these matches
actually had, and it lets a miss be examined for whether the official
explains it.

Nothing here touches a prediction, a probability or a price.
"""

from __future__ import annotations

import argparse
import os
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import check_referees as cr
import polite
import teamnames

BOOK = os.path.join(os.path.dirname(HERE), "forward_test_2026-09-18.xlsx")
COLUMN = "Referee (actual)"
PRETTY = PatternFill("solid", fgColor="DDEBF7")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--url", action="append", default=[],
                    help="an EFL appointments article; repeatable. Defaults "
                         "to the one in referee_appointments.json")
    ap.add_argument("--file", default=BOOK)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()

    urls = a.url
    if not urls:
        st = cr.load_state()
        if not st.get("url"):
            print("no article recorded and none given")
            return 1
        urls = [st["url"]]

    appointed = {}
    for u in urls:
        fx = [f for f in cr.parse_article(polite.get(u)) if f["division"]]
        print(f"{len(fx):>4} league fixtures from {u.rstrip('/').split('/')[-1]}")
        for f in fx:
            appointed[(f["home"], f["away"])] = f["referee"]
    if not appointed:
        print("nothing parsed -- refusing to write")
        return 1

    wb = load_workbook(a.file)
    ws = wb[a.sheet]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    for need in ("Home", "Away", "Referee"):
        if need not in h:
            print(f"sheet {a.sheet!r} has no {need!r} column")
            return 1

    col = h.get(COLUMN)
    if col is None:
        col = ws.max_column + 1
        cell = ws.cell(1, col, COLUMN)
        cell.font = Font("Arial", bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor="1F3864")
        ws.column_dimensions[
            __import__("openpyxl").utils.get_column_letter(col)].width = 17
        print(f"added column {COLUMN!r} at position {col}")

    names = sorted({ws.cell(r, h["Home"]).value for r in range(2, ws.max_row + 1)
                    if ws.cell(r, h["Home"]).value}
                   | {ws.cell(r, h["Away"]).value for r in range(2, ws.max_row + 1)
                      if ws.cell(r, h["Away"]).value})
    resolved = {}
    for (hm, aw), ref in appointed.items():
        h2, a2 = teamnames.resolve(hm, names), teamnames.resolve(aw, names)
        if h2 and a2:
            resolved[(h2, a2)] = ref

    filled = already = missing = 0
    for r in range(2, ws.max_row + 1):
        hm = ws.cell(r, h["Home"]).value
        aw = ws.cell(r, h["Away"]).value
        if not hm or not aw:
            continue
        ref = resolved.get((hm, aw))
        if ref is None:
            missing += 1
            continue
        knew = ws.cell(r, h["Referee"]).value
        if a.write:
            c = ws.cell(r, col, ref)
            c.font = Font("Arial", size=9)
            c.fill = PRETTY
        filled += 1
        tag = "" if str(knew).strip().lower() in ("", "none", "not known") \
            else f"   (model used {knew})"
        print(f"  {'set   ' if a.write else 'would '} {hm} v {aw}: {ref}{tag}")

    print(f"\n{filled} recorded, {missing} with no appointment found")
    if a.write:
        wb.save(a.file)
        print(f"saved {a.file}")
        print("the 'Referee' column was not touched -- it still records what "
              "the model knew")
    else:
        print("dry run -- pass --write to save")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
